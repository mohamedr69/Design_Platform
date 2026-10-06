"""The project's Drawings Review tab: the fire alarm IFC drawing's rooms
looked at by the model against the company's coverage rules (app.review).

  GET  /projects/{id}/drawing-review                          the drawings in force, the rules
  GET  /projects/{id}/drawing-review/{did}                    one drawing's review: floors, rooms, findings
  POST /projects/{id}/drawing-review/{did}/jobs               review it (the IFC worker): HTTP 202 and the job
  POST /projects/{id}/drawing-review/{did}/decisions          accept or dismiss a finding
  GET  /projects/{id}/drawing-review/{did}/image              a piece of the plotted drawing (PNG)
  GET  /projects/{id}/drawing-review/{did}/export.xlsx        the findings as a workbook
  GET  /projects/{id}/drawing-review/{did}/markup.pdf         the accepted changes, floor by floor: the draftsman's schedule
"""
from __future__ import annotations

import io
import re

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import Response, StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from app.core.timeutils import utc_now
from app.database import get_db
from app.deps import get_current_user, require_role
from app.models import ProjectIfcDrawing, User
from app.review import rulings as R
from app.review import service
from app.routers.projects import CREATOR_ROLES, _get_project_or_404
from app.services import activity, jobs

router = APIRouter(tags=["drawing-review"])
KIND = "fa_drawing_review"
_SAFE = re.compile(r"[^\w .()-]+")


def _drawing(db: Session, project, drawing_id: int) -> ProjectIfcDrawing:
    d = db.get(ProjectIfcDrawing, drawing_id)
    if d is None or d.project_id != project.id or d.deleted_at is not None:
        raise HTTPException(404, "Drawing not found")
    return d


@router.get("/projects/{project_id}/drawing-review")
def overview(project_id: int, _current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    from app.ifc.services import revisions
    from app.review import ai

    project = _get_project_or_404(db, project_id)
    drawings = [{"id": d.id, "filename": d.filename, "revision": d.revision or "R0",
                 "pages": service.pages_of(db, project, d)} for d in revisions.in_force(db, project.id)]
    return {"drawings": drawings, "rules": [{"system": k, "name": ai.SYSTEMS[k], "rule": t} for k, t in ai.RULES]}


@router.get("/projects/{project_id}/drawing-review/{drawing_id}")
def get_review(project_id: int, drawing_id: int, _current_user: User = Depends(get_current_user),
               db: Session = Depends(get_db)):
    project = _get_project_or_404(db, project_id)
    out = service.build(db, project, _drawing(db, project, drawing_id))
    db.commit()
    return out


class StartIn(BaseModel):
    # The pages (sheets) to review; none for every floor plan.
    pages: list[int] | None = Field(default=None, max_length=200)


@router.post("/projects/{project_id}/drawing-review/{drawing_id}/jobs", status_code=status.HTTP_202_ACCEPTED)
def start(project_id: int, drawing_id: int, body: StartIn | None = None,
          current_user: User = Depends(require_role(*CREATOR_ROLES)), db: Session = Depends(get_db)):
    from app.routers import jobs as jobs_router
    from app.routers.ifc_boq import _queue_note, _run_inline, _started

    project = _get_project_or_404(db, project_id)
    drawing = _drawing(db, project, drawing_id)
    pages = sorted(set(body.pages)) if body and body.pages else None
    key = f"{KIND}:{project.id}:{drawing.id}"
    existing = jobs.active_by_key(db, key)
    if existing is not None:
        return _started(db, existing, False)
    job, created = jobs.enqueue(db, kind=KIND, project_id=project.id, user_id=current_user.id, dedup_key=key,
                                params={"drawing_id": drawing.id, "pages": pages, "user_id": current_user.id},
                                progress=_queue_note(db), message="")
    if created:
        activity.record(db, current_user, "drawing_review.started",
                        f"Started the drawings review of {drawing.filename} {drawing.revision or ''}".strip(),
                        project=project, entity_type="ifc_drawing", entity_id=drawing.id, detail={"pages": pages})
    if created and jobs_router.RUN_INLINE:
        _run_inline(job.id)
        db.expire_all()
        job = db.get(type(job), job.id)
    return _started(db, job, created)


class DecisionIn(BaseModel):
    id: str = Field(max_length=40)
    # accepted (an issue to fix) | dismissed (not an issue) | open (undo)
    status: str
    note: str = Field(default="", max_length=500)
    # The draftsman's instruction as the engineer words it, over the model's.
    instruction: str = Field(default="", max_length=300)


@router.post("/projects/{project_id}/drawing-review/{drawing_id}/decisions")
def decide(project_id: int, drawing_id: int, body: DecisionIn,
           current_user: User = Depends(require_role(*CREATOR_ROLES)), db: Session = Depends(get_db)):
    project = _get_project_or_404(db, project_id)
    drawing = _drawing(db, project, drawing_id)
    view = service.build(db, project, drawing)
    finding = next((f for f in view["findings"] if f["id"] == body.id), None)
    if finding is None:
        raise HTTPException(404, "That finding is not in the review")
    if body.status not in ("accepted", "dismissed", "open"):
        raise HTTPException(422, "status is accepted, dismissed or open")
    if body.status == "dismissed" and not body.note.strip():
        raise HTTPException(422, "Say why it is not an issue")
    row = service.state(db, project, drawing.id)
    decisions = dict(row.decisions or {})
    # The same comment on the other floors goes with it: those still open,
    # settled only by a ruling, or carried over from another floor -- never
    # one the engineer settled there themselves, or reopened (platform owner,
    # 2 October 2026).
    def carried(f: dict) -> bool:
        own = decisions.get(f["id"])
        if own is not None and f["previous_decision"] is None:
            return bool(own.get("via"))
        return f["decision"] == "open" or f["by_ruling"]

    same = [f for f in service.same_elsewhere(view["findings"], finding) if carried(f)]
    if body.status == "open":
        # reopened: open again, and no ruling settles it until the engineer says
        service.record_decision(decisions, body.id, {"status": "reopened", "by": current_user.id,
                                                    "at": utc_now().isoformat()})
        R.forget(db, project.id, drawing.id, body.id)
        same = [f for f in same if (decisions.get(f["id"]) or {}).get("via") == body.id]
        for f in same:
            decisions.pop(f["id"])
    else:
        now = utc_now().isoformat()
        service.record_decision(decisions, body.id, {"status": body.status, "note": body.note.strip(),
                                                    "instruction": body.instruction.strip(), "by": current_user.id,
                                                    "at": now}, finding)
        for f in same:
            service.record_decision(decisions, f["id"], {
                "status": body.status, "via": body.id,
                "note": f"As on {finding['floor']}: {body.note.strip()}" if body.note.strip() else f"As on {finding['floor']}",
                "instruction": body.instruction.strip(), "by": current_user.id, "at": now}, f)
        # the engineer's word goes to the rules the next review is given
        if finding["action"] != "none":
            R.record(db, project.id, drawing.id, finding, body.status, note=body.note.strip(),
                     instruction=body.instruction.strip(), user_id=current_user.id)
    row.decisions = decisions
    db.commit()
    out = service.build(db, project, drawing)
    out["applied"] = {"count": len(same), "floors": sorted({f["floor"] for f in same})}
    return out


@router.get("/projects/{project_id}/drawing-review/{drawing_id}/image")
def image(project_id: int, drawing_id: int, page: int, box: str = Query(max_length=80), width: int = 900,
          mark: str | None = Query(default=None, max_length=40), _current_user: User = Depends(get_current_user),
          db: Session = Depends(get_db)):
    project = _get_project_or_404(db, project_id)
    drawing = _drawing(db, project, drawing_id)
    try:
        values = [float(v) for v in box.split(",")]
        spot = tuple(float(v) for v in mark.split(",")) if mark else None
        if len(values) != 4 or (spot is not None and len(spot) != 2):
            raise ValueError
    except ValueError as exc:
        raise HTTPException(422, "box is x0,y0,x1,y1 and mark is x,y") from exc
    try:
        png = service.image(project, service.state(db, project, drawing.id), page, tuple(values),
                            max(200, min(width, 2400)), spot)
    except ValueError as exc:
        raise HTTPException(404, str(exc)) from exc
    return Response(png, media_type="image/png", headers={"Cache-Control": "private, max-age=3600"})


@router.get("/projects/{project_id}/drawing-review/{drawing_id}/export.xlsx")
def export(project_id: int, drawing_id: int, _current_user: User = Depends(get_current_user),
           db: Session = Depends(get_db)):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    project = _get_project_or_404(db, project_id)
    drawing = _drawing(db, project, drawing_id)
    view = service.build(db, project, drawing)
    wb = Workbook()
    ws = wb.active
    ws.title = "Findings"
    ws.append([f"EP-{project.ep_number} {project.project_name or ''} - Drawings Review - "
               f"{drawing.filename} {drawing.revision or ''}"])
    ws.append([f"Reviewed by {view['model']}; {view['counts']['reviewed']} of {view['counts']['rooms']} rooms; "
               f"findings are the model's, for the engineer to accept or dismiss"])
    ws.append([])
    ws.append(["#", "Floor", "Sheet", "Room / where", "System", "Action", "Device", "Instruction for the draftsman",
               "Finding", "Seen", "Engineer", "Note"])
    for i, f in enumerate(view["findings"], 1):
        ws.append([i, f["floor"], f["sheet"], f["room"], f["system_name"],
                   {"add": "ADD", "remove": "REMOVE", "replace": "REPLACE"}.get(f["action"], "Unclear"), f["device"],
                   f["instruction"], f["issue"], f["seen"],
                   {"open": "Open", "accepted": "Accepted", "dismissed": "Dismissed"}[f["decision"]], f["note"]])
    for c in ws[4]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="B91C1C")
    ws["A1"].font = Font(bold=True, size=13)
    for i, w in enumerate((5, 22, 10, 30, 28, 10, 26, 60, 45, 26, 12, 30), 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in ws.iter_rows(min_row=5):
        for c in r:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    rooms = wb.create_sheet("Rooms")
    systems = list(view["systems"])
    rooms.append(["Floor", "Room", "What it is"] + [view["systems"][s] for s in systems])
    for f in view["floors"]:
        for r in f["rooms"]:
            rooms.append([f["floor"], r["name"], r["room_type"]]
                         + [((r["checks"] or {}).get(s) or {}).get("status", "not reviewed").replace("_", " ")
                            for s in systems])
    for c in rooms[1]:
        c.font = Font(bold=True)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = _SAFE.sub("_", f"EP-{project.ep_number} Drawings Review.xlsx")
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{name}"'})


@router.get("/projects/{project_id}/drawing-review/{drawing_id}/markup.pdf")
def markup(project_id: int, drawing_id: int, _current_user: User = Depends(get_current_user),
           db: Session = Depends(get_db)):
    """The draftsman's schedule: the accepted changes, floor by floor --
    action, device, room, instruction, and the piece of the plan it is on."""
    from app.ifc import storage
    from app.review.markup import build

    project = _get_project_or_404(db, project_id)
    drawing = _drawing(db, project, drawing_id)
    view = service.build(db, project, drawing)
    if not any(f["decision"] == "accepted" and f["action"] != "none" for f in view["findings"]):
        raise HTTPException(409, "Accept the changes for the draftsman first: the schedule lists accepted changes only")
    row = service.state(db, project, drawing.id)
    doc = build(view, project, str(storage.absolute(row.pdf_path)) if row.pdf_path else None)
    pdf = doc.tobytes(garbage=3, deflate=True)
    doc.close()
    name = _SAFE.sub("_", f"EP-{project.ep_number} Draftsman Schedule {drawing.revision or ''}.pdf")
    return Response(pdf, media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="{name}"'})


class BulkIn(BaseModel):
    ids: list[str] = Field(max_length=2000)
    # accepted | open
    status: str


@router.post("/projects/{project_id}/drawing-review/{drawing_id}/decisions/bulk")
def decide_many(project_id: int, drawing_id: int, body: BulkIn,
                current_user: User = Depends(require_role(*CREATOR_ROLES)), db: Session = Depends(get_db)):
    """Accept (or reopen) many findings at once -- the ones on show, after
    the engineer has looked through them. Dismissing needs a reason each."""
    project = _get_project_or_404(db, project_id)
    drawing = _drawing(db, project, drawing_id)
    if body.status not in ("accepted", "open"):
        raise HTTPException(422, "status is accepted or open")
    known = {f["id"]: f for f in service.build(db, project, drawing)["findings"]}
    row = service.state(db, project, drawing.id)
    decisions = dict(row.decisions or {})
    stamp = utc_now().isoformat()
    for fid in body.ids:
        f = known.get(fid)
        if f is None or f["action"] == "none":
            continue
        if body.status == "open":
            service.record_decision(decisions, fid, {"status": "reopened", "by": current_user.id, "at": stamp})
            R.forget(db, project.id, drawing.id, fid)
        elif f["decision"] != "accepted":
            # not accepted as it stands now -- never decided, or decided on another proposal
            service.record_decision(decisions, fid, {"status": "accepted", "note": "", "instruction": "",
                                                     "by": current_user.id, "at": stamp}, f)
            R.record(db, project.id, drawing.id, f, "accepted", user_id=current_user.id)
    row.decisions = decisions
    db.commit()
    activity.record(db, current_user, "drawing_review.bulk", f"{body.status.title()} {len(body.ids)} drawings-review findings",
                    project=project, entity_type="ifc_drawing", entity_id=drawing.id)
    return service.build(db, project, drawing)


@router.delete("/drawing-review/rulings/{ruling_id}")
def delete_ruling(ruling_id: int, current_user: User = Depends(require_role(*CREATOR_ROLES)),
                  db: Session = Depends(get_db)):
    """Take a ruling out of what the review is given (the finding it came
    from keeps the engineer's decision)."""
    from app.models import ReviewRuling

    ruling = db.get(ReviewRuling, ruling_id)
    if ruling is None:
        raise HTTPException(404, "No such ruling")
    db.delete(ruling)
    db.commit()
    activity.record(db, current_user, "drawing_review.ruling_deleted",
                    f"Removed the review ruling: {ruling.decision} {ruling.action} {ruling.device} in {ruling.room}")
    return {"deleted": ruling_id}
