"""The Fire Alarm Interface Schedule as a workbook, in the order the
engineer reads it: which drawings were received, the building's floors,
the schedule floor by floor, the floor and equipment summaries, the totals,
what is left to verify, the conflicts, and the interface matrix applied.

Every total is summed from the schedule's own lines (`service.build`), so
the summaries always agree with the schedule."""
from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
FLOOR_FILL = PatternFill("solid", fgColor="DCE6F1")
TOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")
STATUS_FILL = {"available": "C6EFCE", "missing": "FFC7CE", "failed": "FFC7CE", "not_provided": "EDEDED"}
THIN = Border(bottom=Side(style="thin", color="D9D9D9"))

SCHEDULE_HEADERS = ["S.No", "Floor", "Location", "Equipment Tag", "Equipment Type", "Equipment Description", "System",
                    "Required Contact", "Monitoring Signals", "Control Signals", "Alarm / Monitoring Function",
                    "Supervisory Function", "Required Third-Party Action", "CT1", "CT2", "CR",
                    "Estimated FA Module Qty", "Source Drawing", "Drawing Ref", "Confidence", "Evidence"]
SCHEDULE_WIDTHS = [6, 20, 18, 18, 26, 42, 18, 16, 11, 11, 40, 40, 40, 6, 6, 6, 12, 40, 36, 16, 34]


def evidence_note(view: dict) -> str:
    """What the schedule's evidence is, said on every sheet and page (FI-P1 r3 S10)."""
    primary, state = view.get("primary", "current"), view.get("view_state", "current")
    reasons = "; ".join(view.get("view_reasons") or [])
    if primary == "published":
        at = (view.get("published_at") or "")[:16].replace("T", " ")
        return f"LAST PUBLISHED {at}, NOT VERIFIED NOW" + (f" ({reasons})" if reasons else "")
    if primary == "none":
        return "NOT READ COMPLETELY: no schedule" + (f" ({reasons})" if reasons else "")
    if state != "current":
        return "PROVISIONAL" + (f" ({reasons})" if reasons else "")
    return "Current: every drawing verified now"


def workbook(view: dict) -> io.BytesIO:
    wb = Workbook()
    wb.remove(wb.active)
    p = view["project"]
    title = f"EP-{p['ep_number']} {p.get('name') or ''} - Fire Alarm Interface Schedule".strip()
    scanned = (view.get("scanned_at") or "")[:16].replace("T", " ") or "not read yet"
    info = (f"{evidence_note(view)} | Drawings read {scanned}; exported {datetime.now():%Y-%m-%d %H:%M}; "
            f"interface matrix: {view['matrix']['name']}")
    known = view.get("totals_known", True)

    def num(value):
        return value if known else "—"

    def sheet(name: str, headers: list[str], widths: list[int], *, heading: str | None = None):
        ws = wb.create_sheet(name[:31])
        ws.append([heading or title])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([info])
        ws.append([])
        ws.append(headers)
        for c in ws[4]:
            c.font, c.fill = HEAD_FONT, HEAD_FILL
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A5"
        return ws

    def band(ws, fill, bold=True):
        for c in ws[ws.max_row]:
            c.fill = fill
            c.font = Font(bold=bold)

    def wrap(ws):
        for r in ws.iter_rows(min_row=5):
            for c in r:
                c.alignment = Alignment(vertical="top", wrap_text=True)

    # A. Drawing coverage
    ws = sheet("A. Drawing Coverage", ["Discipline", "Status", "Folder", "Drawing", "Revision", "Modified", "Sheets",
                                       "Floor plans", "Items found", "Remarks"],
               [24, 14, 40, 44, 9, 17, 8, 11, 11, 60])
    for d in view["coverage"]:
        status = d.get("badge_text") or {"available": "Available", "missing": "Missing", "failed": "Could not be read",
                                         "not_provided": "Not provided"}.get(d["status"], d["status"])
        if not d["files"]:
            ws.append([d["name"], status, d["folder"] or "-", "-", "", "", "", "", "", d["purpose"]])
            ws.cell(ws.max_row, 2).fill = PatternFill("solid", fgColor=STATUS_FILL.get(d["status"], "FFFFFF"))
            continue
        for f in d["files"]:
            remarks = f.get("error") or ("Superseded by a later revision in the folder: not read"
                                         if f["status"] == "superseded" else
                                         "Received, not read yet: read the drawings again"
                                         if f["status"] == "unread" else
                                         f"Last known only, not counted ({f.get('reason') or 'not current'})"
                                         if f["status"] == "stale" else
                                         "Not readable by the schedule (PDF or other): not counted"
                                         if f["status"] == "unsupported" else d["purpose"])
            if f.get("lifts"):
                remarks += f"; {f['lifts']} lift labels"
            ws.append([d["name"], status, d["folder"] or "Fire alarm IFC drawing in force", f["filename"],
                       f.get("revision") or "", (f.get("modified") or "")[:16].replace("T", " "), f["sheets"], f["plans"],
                       f["items"], remarks])
            ws.cell(ws.max_row, 2).fill = PatternFill("solid", fgColor=STATUS_FILL.get(d["status"], "FFFFFF"))
    wrap(ws)

    # B. Building floors
    ws = sheet("B. Building Floors", ["#", "Floor", "Key", "Named on", "Interface lines", "Monitoring", "Control"],
               [5, 24, 12, 50, 14, 12, 12])
    by_floor = {f["floor_key"]: f for f in view["floor_summary"]}
    for i, f in enumerate(view["floors"], 1):
        s = by_floor.get(f["key"], {})
        named = f.get("ifc_sheet") or ""
        if f.get("registered") is False:
            named = f"Not in the Building Floor Registry -- named by {named}"
        ws.append([i, f["name"], f["key"], named, s.get("items", 0), s.get("monitoring", 0), s.get("control", 0)])
    listed = {f["key"] for f in view["floors"]}
    for s in view["floor_summary"]:
        if s["floor_key"] not in listed:
            ws.append(["", s["floor"], s["floor_key"], "Not in the Building Floor Registry: see H. Conflicts",
                       s["items"], s["monitoring"], s["control"]])
    wrap(ws)

    # C. The schedule, floor by floor
    ws = sheet("C. Interface Schedule", SCHEDULE_HEADERS, SCHEDULE_WIDTHS)
    rows = view["rows"]
    current = None
    floor_rows: list[dict] = []

    def floor_total():
        if not floor_rows:
            return
        m = sum(r["monitoring"] for r in floor_rows)
        c = sum(r["control"] for r in floor_rows)
        mods = {k: sum(r["modules"][k] for r in floor_rows) for k in ("CT1", "CT2", "CR")}
        ws.append(["", f"{current} total", "", "", f"{len(floor_rows)} line{'s' if len(floor_rows) != 1 else ''}", "", "", "", m, c, "", "", "",
                   mods["CT1"], mods["CT2"], mods["CR"], sum(mods.values())])
        band(ws, TOTAL_FILL)

    for r in rows:
        if r["floor"] != current:
            floor_total()
            current, floor_rows = r["floor"], []
            ws.append([f"{r['floor']}"])
            band(ws, FLOOR_FILL)
        floor_rows.append(r)
        ws.append([r["no"], r["floor"], r["location"] or "-", r["tag"], r["equipment"], r["description"], r["system"],
                   r["contacts"], r["monitoring"], r["control"], r["alarm"] or "-", r["supervisory"] or "-",
                   r["action"] or "-", r["modules"]["CT1"], r["modules"]["CT2"], r["modules"]["CR"], r["module_qty"],
                   r["source"], r["drawing_ref"], r["confidence"], r["evidence"]])
    floor_total()
    t = view["totals"]
    if rows:
        ws.append(["", "BUILDING TOTAL", "", "", f"{num(t['items'])} lines", "", "", "", num(t["monitoring"]),
                   num(t["control"]), "", "", "", num(t["modules"]["CT1"]), num(t["modules"]["CT2"]),
                   num(t["modules"]["CR"]), num(t["module_qty"])])
        band(ws, HEAD_FILL)
        for c in ws[ws.max_row]:
            c.font = HEAD_FONT
    else:
        ws.append(["", "Nothing scheduled yet: read the drawings, then settle Verification Required."])
    wrap(ws)

    # D. Floor summary
    ws = sheet("D. Floor Summary", ["Floor", "Interface lines", "Monitoring Signals", "Control Signals",
                                    "Total Interfaces", "CT1", "CT2", "CR", "FA Modules"],
               [24, 14, 16, 14, 15, 7, 7, 7, 12])
    for s in view["floor_summary"]:
        ws.append([s["floor"], s["items"], s["monitoring"], s["control"], s["total"], s["modules"]["CT1"],
                   s["modules"]["CT2"], s["modules"]["CR"], s["module_qty"]])
    ws.append(["Total", t["items"], t["monitoring"], t["control"], t["interface_points"], t["modules"]["CT1"],
               t["modules"]["CT2"], t["modules"]["CR"], t["module_qty"]])
    band(ws, TOTAL_FILL)

    # E. Equipment-type summary
    ws = sheet("E. Equipment Summary", ["Equipment Type", "Required Contact", "Qty", "Monitoring Signals",
                                        "Control Signals", "CT1", "CT2", "CR", "FA Modules"],
               [32, 22, 8, 16, 14, 7, 7, 7, 12])
    for s in view["type_summary"]:
        ws.append([s["equipment"], s["contacts"], s["qty"], s["monitoring"], s["control"], s["modules"]["CT1"],
                   s["modules"]["CT2"], s["modules"]["CR"], s["module_qty"]])
    ws.append(["Total", "", t["items"], t["monitoring"], t["control"], t["modules"]["CT1"], t["modules"]["CT2"],
               t["modules"]["CR"], t["module_qty"]])
    band(ws, TOTAL_FILL)

    # F. Overall totals
    ws = sheet("F. Totals", ["Item", "Value", "Basis"], [40, 12, 70])
    for label, value, basis in (
        ("Interface lines", num(t["items"]), "One line per item per floor"),
        ("Total monitoring signals", num(t["monitoring"]), "Sum of the lines' monitoring signals"),
        ("Total control signals", num(t["control"]), "Sum of the lines' control signals"),
        ("Total interface points", num(t["interface_points"]), "Monitoring + control"),
        ("Monitoring modules (CT1 + CT2)", num(t["monitoring_modules"]), "From the matrix's Required Contact: CT2 is one dual-input module"),
        ("  of which CT1", num(t["modules"]["CT1"]), "Single-input monitor module"),
        ("  of which CT2", num(t["modules"]["CT2"]), "Dual-input monitor module"),
        ("Control modules (CR)", num(t["control_modules"]), "Relay module; a lift takes two (2NO.CR FOR EACH LIFT)"),
        ("Total estimated FA modules", num(t["module_qty"]), "Estimate: confirm against the FA panel's module configuration"),
        ("Items still to verify (not counted)", num(t["to_verify"]), "G. Verification Required"),
    ):
        ws.append([label, value, basis])

    # G. Verification required
    ws = sheet("G. Verification Required", ["#", "Floor (proposed)", "Possible Equipment", "System", "Issue",
                                            "Source Drawing", "Evidence", "Proposed Qty", "Required Verification"],
               [5, 22, 28, 18, 60, 40, 60, 12, 44])
    for i, g in enumerate(view["verification"], 1):
        ws.append([i, ", ".join(g["proposed_floors"]) or "Not identified", g["equipment"], g["system"], g["reason"],
                   g["source"], g["evidence"], g["proposed_qty"] if g["proposed_qty"] is not None else "-",
                   "Confirm the floor and quantity (and tags), or say it is not an interface"])
    if not view["verification"]:
        ws.append(["", "Nothing left to verify."])
    wrap(ws)

    # H. Conflicts and missing information
    ws = sheet("H. Conflicts & Missing", ["#", "Kind", "Detail"], [5, 28, 120])
    n = 0
    for d in view["coverage"]:
        if d["status"] in ("missing", "failed") or not d.get("read", True):
            n += 1
            ws.append([n, "Drawings missing" if d["status"] == "missing" else "Drawings not read",
                       f"{d['name']} ({d['folder']}): the {d['purpose'].lower()} are not in this schedule"])
    for c in view["conflicts"]:
        n += 1
        ws.append([n, "Conflict", c])
    for e in view["excluded_found"]:
        n += 1
        ws.append([n, "Excluded system (not scheduled)", f"{e['equipment']} on {e['sheet']} of {e['source']}: \"{e['text']}\""])
    for r in view["rejected"]:
        n += 1
        ws.append([n, "Not an interface (engineer)", f"{r['equipment']} {r['tag']} on {r['floor']} ({r['source']}): {r.get('reason', '')}"])
    for g in view["settled"]:
        if g["status"] == "dismissed":
            n += 1
            ws.append([n, "Not scheduled (engineer)", f"{g['equipment']} ({g['source']}): {(g.get('decision') or {}).get('reason', '')}"])
    if view["matrix"]["unclear_rows"]:
        n += 1
        ws.append([n, "Interface matrix", f"Rows {', '.join(map(str, view['matrix']['unclear_rows']))} of the matrix are "
                                          "not legible on the copy transcribed: they are not applied."])
    if not n:
        ws.append(["", "None"])
    wrap(ws)

    # The matrix applied
    ws = sheet("Interface Matrix", ["S.No", "Equipment", "Required Contacts", "No. of Monitoring Signals",
                                    "No. of Control Signals", "Alarm", "Supervisory", "Action for Third Party Equipment",
                                    "Scheduled"],
               [6, 30, 22, 12, 12, 50, 50, 50, 22], heading=view["matrix"]["name"])
    for r in view["matrix"]["rules"]:
        ws.append([r["no"], r["name"], r["contacts"], r["monitoring"] or "-", r["control"] or "-", r["alarm"] or "-",
                   r["supervisory"] or "-", r["action"] or "-", "Excluded" if r["excluded"] else "Yes"])
    wrap(ws)

    # Last known, not current: kept for audit, never counted (FI-P1 r3 S10)
    if view.get("last_known"):
        ws = sheet("Last known, not current", ["Package", "Drawing", "Why not current", "Last read", "Labels by type"],
                   [18, 50, 30, 18, 70],
                   heading=f"{title} -- last known readings, NOT counted in this schedule")
        for k in view["last_known"]:
            ws.append([k.get("package"), k.get("filename") or k.get("relative_path"), k.get("reason") or "",
                       (k.get("last_known_at") or "")[:16].replace("T", " "),
                       ", ".join(f"{key} {n}" for key, n in sorted((k.get("counts_by_key") or {}).items()))])
        wrap(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
