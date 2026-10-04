"""The mechanical equipment schedules -- Schedule of Fans, Schedule of FAHU,
... -- read to check the interface schedule's fans, FAHUs and AHUs (§35 of
the interface schedule's instructions).

A schedule is a table: a row per item, or per run of items ("B3-SEF-1 TO
4", 4 Nos, BASEMENT-3). The tag column is found by its heading (TAG, REF,
REFERENCE), the location and quantity columns likewise; a run is expanded
into its tags. Each tag is read as matrix equipment the way a drawing's tag
is (app.interfaces.detect) -- a chiller or a sump pump the matrix has no row
for is counted and left aside.

What a schedule is used for (app.interfaces.service): a tag the plans place
is confirmed by it; a tag the plans do not place is asked about, never
scheduled on the schedule's word alone; a floor-qualified tag ("B2-SEF-1")
tells the plans' repeated "SEF-1" on each basement apart.
"""
from __future__ import annotations

import os
import re
from pathlib import Path

from app.interfaces import detect
from app.interfaces.matrix import BY_KEY
from app.services import document_control

SCHEDULE_VERSION = "1"
SUFFIXES = (".xlsx", ".xlsm", ".xls")
_TAG_HEAD = re.compile(r"\b(?:TAG|REF(?:ERENCE)?\.?|EQUIPMENT\s*(?:NO|ID)|UNIT\s*NO)\b", re.I)
_LOCATION_HEAD = re.compile(r"\b(?:LOCATION|FLOOR|LEVEL)\b", re.I)
_QTY_HEAD = re.compile(r"\b(?:QTY|QUANTITY|NOS|NO'S)\b", re.I)      # never "SR. NO."
_RANGE = re.compile(r"^(.*?)(\d+)\s*(?:TO|~|–|-)\s*(\d+)\s*$", re.I)
# A tag that names its floor first: "B3-SEF-1" is SEF-1 of the 3rd basement.
FLOOR_PREFIX = re.compile(r"^(B\d{1,2}|P\d{1,2}|L\d{1,3}|GF|RF|TRF|MF)\s*-\s*(.+)$", re.I)
_MAX_ROWS = 2000


def _rows(path: Path) -> list[tuple[str, list[list[str]]]]:
    """(sheet name, its rows as text) of a workbook."""
    out = []
    if path.suffix.lower() == ".xls":
        import xlrd

        book = xlrd.open_workbook(document_control._os_path(path))
        for sheet in book.sheets():
            out.append((sheet.name, [[_cell(v) for v in sheet.row_values(i)] for i in range(min(sheet.nrows, _MAX_ROWS))]))
        return out
    from openpyxl import load_workbook

    book = load_workbook(document_control._os_path(path), data_only=True, read_only=True)
    try:
        for ws in book.worksheets:
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= _MAX_ROWS:
                    break
                rows.append([_cell(v) for v in row])
            out.append((ws.title, rows))
    finally:
        book.close()
    return out


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return " ".join(str(v).split())


def expand(tag: str) -> list[str]:
    """A schedule's tag cell as tags: "B3-SEF-1 TO 4" -> B3-SEF-1 .. B3-SEF-4,
    "CHILLER- 01 TO 04" -> CHILLER-01 .. CHILLER-04 (the width kept)."""
    t = " ".join(tag.split()).upper()
    m = _RANGE.match(t)
    if m and m.group(1) and re.search(r"[A-Z]", m.group(1)):
        a, b = int(m.group(2)), int(m.group(3))
        if 0 <= a < b <= a + 200:
            width = len(m.group(2)) if m.group(2).startswith("0") else 0
            stem = m.group(1).replace(" ", "")
            return [f"{stem}{str(n).zfill(width)}" for n in range(a, b + 1)]
    return [t]


def base_tag(tag: str) -> tuple[str | None, str]:
    """(the floor a tag names first, the tag without it): "B3-SEF-1" ->
    ("B3", "SEF-1"); "FAHU-01" -> (None, "FAHU-01")."""
    m = FLOOR_PREFIX.match(tag.strip())
    return (m.group(1).upper(), m.group(2).upper()) if m else (None, tag.strip().upper())


def same_tag(a: str, b: str) -> bool:
    """FAHU-01 and FAHU-1 are one tag: numbers compared as numbers."""
    def key(t: str) -> str:
        return re.sub(r"\d+", lambda m: str(int(m.group())), detect.tag_key(t))
    return key(a) == key(b)


def tag_identity(tag: str) -> str:
    return re.sub(r"\d+", lambda m: str(int(m.group())), detect.tag_key(tag))


def read(path: Path) -> dict:
    """Every row of every sheet that names matrix equipment: its tags, the
    location as written, the quantity, and where in the workbook it is."""
    items: list[dict] = []
    other = 0
    for sheet, rows in _rows(path):
        header_at, cols = None, {}
        for i, row in enumerate(rows[:40]):
            tag_col = next((j for j, v in enumerate(row) if v and _TAG_HEAD.search(v) and len(v) < 40), None)
            if tag_col is not None:
                header_at = i
                cols["tag"] = tag_col
                cols["location"] = next((j for j, v in enumerate(row) if v and _LOCATION_HEAD.search(v)), None)
                cols["qty"] = next((j for j, v in enumerate(row) if v and _QTY_HEAD.search(v)), None)
                break
        if header_at is None:
            continue
        title = " ".join(v for r in rows[:header_at] for v in r if v)
        for n, row in enumerate(rows[header_at + 1:], start=header_at + 2):
            cell = row[cols["tag"]] if cols["tag"] < len(row) else ""
            if not cell or len(cell) > 60:
                continue
            words = " ".join(v for v in row if v)
            # "B3-SEF-1 TO 4" is read as "SEF-1 TO 4" on the 3rd basement
            _floor, bare = base_tag(cell)
            found = detect.detect(bare, set(BY_KEY)) or []
            if not found:
                # the tag alone does not say ("CSEF-1 TO 4" under SMOKE MANAGEMENT): the row's words may
                found = [d for d in detect.detect(f"{cell} {words}"[:150], set(BY_KEY)) if d.kind == detect.LABEL]
            found = [d for d in found if not BY_KEY[d.key].excluded]
            if not found:
                other += 1
                continue
            d = found[0]
            # "Pump Room FAF-1": the tag is the one the words name; a run is expanded
            # (a floor-qualified tag keeps its floor: "B2-SEF-1", never "SEF-1")
            whole = _floor is not None or _RANGE.match(" ".join(cell.split()).upper())
            tags = expand(cell) if whole or not d.tag else expand(d.tag)
            location = row[cols["location"]] if cols.get("location") is not None and cols["location"] < len(row) else ""
            qty_text = row[cols["qty"]] if cols.get("qty") is not None and cols["qty"] < len(row) else ""
            qty = int(qty_text) if qty_text.isdigit() else None
            items.append({"key": d.key, "tags": tags, "location": location, "qty": qty, "sheet": sheet, "row": n,
                          "text": cell, "title": title[:120]})
    return {"schedule_version": SCHEDULE_VERSION, "items": items, "other_rows": other}


def discover(root: Path) -> list[Path]:
    """The schedules in the project's mechanical IFC folders: a folder named
    for schedules ("MECHANICAL SCHEDULE") or a workbook beside the drawings."""
    base = root / "03- Drawings" / "IFC" / "Mechanical"
    found = []
    for dirpath, _dirs, names in os.walk(document_control._os_path(base)):
        for name in names:
            if name.lower().endswith(SUFFIXES) and not name.startswith("~$"):
                found.append(Path(dirpath) / name)
    return sorted(found, key=lambda p: p.name.lower())
