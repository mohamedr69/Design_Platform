"""FI-P1 r3 Stage 0.1 -- the interface schedule's evidence: current, stale,
published, and never an empty current schedule from a folder that cannot be
seen. Contract: docs/milestones/fa-interfaces/FI-P1-R2-CORRECTION/CORRECTION-R3.md
with the binding conditions C1-C8 of REVIEW-R3.md. Test ids are the contract's.
"""
from __future__ import annotations

import io
import json
import os
import shutil
from pathlib import Path
from types import SimpleNamespace

import ezdxf
import openpyxl
import pytest

from app.core.config import get_settings
from app.interfaces import evidence, scan, service
from app.models import Project, ProjectFaInterfaces
from app.services import jobs
from tests.conftest import login
from tests.test_fa_interfaces import _fire_fighting, _smoke_management, _ventilation

settings = get_settings()
REAL_READ = scan.read
REAL_SHA = service._sha256


@pytest.fixture
def p(client, db_session, tmp_path):
    """A project with SM, FF and HVAC drawings in its IFC folders."""
    login(client, settings.default_admin_email, settings.default_admin_password)
    root = tmp_path / "EP-40900 Tower"
    ifc = root / "03- Drawings" / "IFC"
    for sub in ("Mechanical/SM", "Mechanical/FF", "Mechanical/HVAC"):
        (ifc / sub).mkdir(parents=True)
    _smoke_management(ifc / "Mechanical/SM" / "SM LAYOUT.dxf")
    _fire_fighting(ifc / "Mechanical/FF" / "FF LAYOUT R1.dxf")
    _ventilation(ifc / "Mechanical/HVAC" / "GROUND FLOOR VENTILATION LAYOUT.dxf")
    pid = client.post("/projects", json={"ep_number": "40900", "project_name": "Tower", "design_sheets": [],
                                         "source_folder_path": str(root)}).json()["id"]
    return SimpleNamespace(client=client, db=db_session, pid=pid, root=root, ifc=ifc)


def _project(p) -> Project:
    p.db.expire_all()
    return p.db.get(Project, p.pid)


def _scan(p, **kw) -> dict:
    out = service.scan_project(p.db, _project(p), **kw)
    p.db.expire_all()
    return out


def _row(p) -> ProjectFaInterfaces:
    p.db.expire_all()
    return p.db.query(ProjectFaInterfaces).filter(ProjectFaInterfaces.project_id == p.pid).one()


def _entry(p, suffix: str) -> dict:
    return next(e for e in _row(p).sources if (e.get("relative_path") or "").endswith(suffix))


def _view(p) -> dict:
    return p.client.get(f"/projects/{p.pid}/fa-interfaces").json()


def _badge(view: dict, code: str) -> dict:
    return next(c for c in view["coverage"] if c["discipline"] == code)


def _rewrite(path: Path, extra: str = "CHANGED") -> None:
    """New bytes, new size and time: a drawing that changed."""
    doc = ezdxf.readfile(path)
    doc.modelspace().add_text(extra, height=10).set_placement((999_999, 999_999))
    doc.saveas(path)


def _failing_read(monkeypatch, fragment: str) -> None:
    def read(path, discipline, check=None):
        if fragment in str(path):
            raise RuntimeError("converter crashed")
        return REAL_READ(path, discipline, check=check)
    monkeypatch.setattr(scan, "read", read)


@pytest.fixture
def complete(p):
    """The first complete read: everything current and published."""
    out = _scan(p)
    assert out["published"] is True
    row = _row(p)
    assert row.published_basis == "complete_scan" and row.generation == 1
    return p


# --- T-01 / T-06: a failed reread, then a successful one ------------------------------------------------


def test_T01_a_failed_reread_is_stale_not_counted_and_the_published_schedule_is_shown(complete, monkeypatch):
    p = complete
    before = _view(p)
    ff_lines = [r for r in before["rows"] if r["key"] == "zone_control_valve"]
    assert ff_lines
    group = before["verification"][0]
    p.client.post(f"/projects/{p.pid}/fa-interfaces/decisions", json={"id": group["id"], "action": "dismiss",
                                                                       "reason": "not here"})
    decisions = dict(_row(p).decisions)
    _rewrite(p.ifc / "Mechanical/FF" / "FF LAYOUT R1.dxf")
    _failing_read(monkeypatch, "FF LAYOUT")
    _scan(p)
    ff = _entry(p, "FF LAYOUT R1.dxf")
    assert ff["status"] == "stale" and ff["stale_reason"] == "read_failed"
    assert ff["last_known"]["result"]["items"] and "result" not in ff          # kept apart, not counted
    assert _row(p).decisions == decisions                                      # never written by a read
    view = _view(p)
    assert view["view_state"] == "provisional" and view["primary"] == "published"
    assert [r["id"] for r in view["rows"] if r["key"] == "zone_control_valve"] == [r["id"] for r in ff_lines]
    assert view["current_summary"]["rows"] < len(view["rows"])                # current excludes FF
    assert _badge(view, "FF")["badge"] == "received_not_read" and _badge(view, "FF")["received"] is True
    assert any(k["relative_path"].endswith("FF LAYOUT R1.dxf") and k["reason"] == "read_failed" for k in view["last_known"])

    # T-06: read again, it reads -> current, the published schedule moves to it
    monkeypatch.setattr(scan, "read", REAL_READ)
    out = _scan(p)
    assert out["published"] is True and _entry(p, "FF LAYOUT R1.dxf")["status"] == "read"
    view = _view(p)
    assert (view["view_state"], view["primary"], view["last_known"]) == ("current", "current", [])


# --- T-02 / C1: a folder or a file missing -------------------------------------------------------------


def test_T02_a_missing_folder_keeps_the_last_reading_apart_and_is_never_received(complete):
    p = complete
    shutil.move(str(p.ifc / "Mechanical/SM"), str(p.ifc / "Mechanical/SM-moved"))
    _scan(p)
    sm = _entry(p, "SM LAYOUT.dxf")
    assert (sm["status"], sm["stale_reason"]) == ("stale", "folder_missing")
    view = _view(p)
    badge = _badge(view, "SM")
    assert badge["badge"] == "missing_last_known" and badge["received"] is False and badge["status"] == "missing"
    assert view["primary"] == "published" and view["view_state"] == "provisional"


def test_C1_a_file_missing_from_a_present_folder_is_not_removed_and_publishes_nothing(complete):
    p = complete
    published = _row(p).published
    (p.ifc / "Mechanical/HVAC" / "GROUND FLOOR VENTILATION LAYOUT.dxf").unlink()
    (p.ifc / "Mechanical/HVAC" / "readme.pdf").write_bytes(b"%PDF-1.4")      # the folder is present
    out = _scan(p)
    hvac = _entry(p, "GROUND FLOOR VENTILATION LAYOUT.dxf")
    assert (hvac["status"], hvac["stale_reason"]) == ("stale", "missing")    # not removed by itself
    assert out["published"] is False and _row(p).published == published
    view = _view(p)
    assert view["primary"] == "published"
    assert any(r["key"] == "ahu" for r in view["rows"])                      # the published AHU still shown


def test_C1_a_folder_listing_brought_down_late_is_listing_failed(complete, monkeypatch):
    p = complete
    real = evidence.attributes
    monkeypatch.setattr(evidence, "attributes",
                        lambda path: evidence.ATTR_RECALL_ON_OPEN if path.rstrip("\\/").endswith("HVAC") else real(path))
    out = _scan(p)
    hvac = _entry(p, "GROUND FLOOR VENTILATION LAYOUT.dxf")
    assert (hvac["status"], hvac["stale_reason"]) == ("stale", "listing_failed")
    assert out["published"] is False
    assert any("HVAC" in d for d in _view(p)["evidence"]["listing_failed"])


def test_T19_a_listing_error_under_a_folder_is_listing_failed_and_blocks_publishing(complete, monkeypatch):
    p = complete
    real_walk = os.walk

    def walk(top, onerror=None, **kw):
        if str(top).endswith("SM"):
            onerror(PermissionError(13, "denied", str(top)))
            return iter(())
        return real_walk(top, onerror=onerror, **kw)
    monkeypatch.setattr(evidence.os, "walk", walk)
    out = _scan(p)
    assert _entry(p, "SM LAYOUT.dxf")["stale_reason"] == "listing_failed" and out["published"] is False


# --- T-03 / T-04 / T-21: the folder cannot be seen ----------------------------------------------------


def test_T03_an_unreachable_folder_fails_the_read_writes_nothing_and_shows_the_published_schedule(complete, monkeypatch):
    p = complete
    row = _row(p)
    sources, generation, published = json.dumps(row.sources, sort_keys=True), row.generation, row.published
    shutil.move(str(p.root), str(p.root) + "-offline")
    with pytest.raises(service.SourceUnreachable):
        _scan(p)
    row = _row(p)
    assert (json.dumps(row.sources, sort_keys=True), row.generation, row.published) == (sources, generation, published)
    view = _view(p)
    assert (view["view_state"], view["primary"], view["current_summary"]) == ("unverified", "published", None)
    assert {k["reason"] for k in view["last_known"]} == {"unreachable"}
    assert view["rows"] and view["totals_known"] is True
    folders = [c for c in view["coverage"] if c["folder"]]
    assert {c["badge"] for c in folders} == {"unreachable"} and not any(c["status"] == "available" for c in folders)
    # through the job: failed, said plainly, nothing written
    import app.routers.jobs as jobs_router

    monkeypatch.setattr(jobs_router, "RUN_INLINE", True)
    started = p.client.post(f"/projects/{p.pid}/fa-interfaces/scan/jobs").json()
    job = p.client.get(f"/jobs/{started['id']}").json()
    assert job["status"] == "failed" and "cannot be reached" in job["error"] and "Nothing was changed" in job["error"]


def test_T04_an_unsynced_drawings_folder_never_becomes_an_empty_schedule(complete):
    p = complete
    (p.ifc / "Mechanical/SM" / "SM SHOP DRAWING.pdf").write_bytes(b"%PDF-1.4")
    _scan(p)
    shutil.move(str(p.ifc), str(p.ifc) + "-cloud")
    out = _scan(p)                                                            # completes, writes stale states
    pdf = _entry(p, "SM SHOP DRAWING.pdf")
    assert pdf["status"] == "unsupported" and pdf["present"] is False         # listed, not removed
    assert out["published"] is False
    assert {e["stale_reason"] for e in _row(p).sources if e["kind"] in ("folder", "schedule")} == {"ifc_root_missing"}
    view = _view(p)
    assert (view["view_state"], view["primary"]) == ("unverified", "published")
    assert view["rows"] and view["totals"]["items"] == len(view["rows"]) > 0
    assert {c["badge"] for c in view["coverage"] if c["folder"]} == {"not_synced"}


def test_T21_a_project_without_a_folder(p, client):
    project = _project(p)
    project.source_folder_path = None
    p.db.commit()
    view = _view(p)
    assert view["view_state"] == "unverified" and view["primary"] == "none" and view["totals_known"] is False
    assert {c["badge"] for c in view["coverage"] if c["folder"]} == {"no_folder"}
    assert client.post(f"/projects/{p.pid}/fa-interfaces/scan/jobs").status_code == 422


# --- T-05 / C8: the fire alarm IFC drawing is not judged by the folder ---------------------------------


def _fa_dxf(path: Path) -> Path:
    doc = ezdxf.new("R2018")
    doc.header["$INSUNITS"] = 4
    doc.modelspace().add_text("LIFT 1", height=200).set_placement((0, 0))
    doc.saveas(path)
    return path


def _fa_in_force(monkeypatch, dxf: Path, *, exists=True):
    rel = "03- Drawings/IFC/Electrical/FA/FIRE ALARM LAYOUT.dwg"
    monkeypatch.setattr(service, "fa_in_force", lambda db, project: {rel: {
        "path": str(dxf), "dxf_exists": exists, "sha256": "f" * 64, "filename": "FIRE ALARM LAYOUT.dwg",
        "revision": "R0", "fa_drawing_id": 1, "size": dxf.stat().st_size if exists else None, "mtime": 1.0}})
    return rel


def test_T05_the_EP30880_shape_an_empty_architectural_folder_and_a_fire_alarm_ifc_in_force(p, monkeypatch, tmp_path):
    (p.ifc / "Architectural").mkdir(exist_ok=True)          # made empty with the project, as on EP-30880
    assert not any((p.ifc / "Architectural").iterdir())
    rel = _fa_in_force(monkeypatch, _fa_dxf(tmp_path / "fa.dxf"))
    out = _scan(p)
    assert out["published"] is True
    fa = _entry(p, "FIRE ALARM LAYOUT.dwg")
    assert fa["kind"] == "fa_ifc" and fa["status"] == "read"
    view = _view(p)
    assert _badge(view, "ARCH")["badge"] == "received" and view["view_state"] == "current"
    assert rel in {s["relative_path"] for s in _row(p).published["sources"]}


def test_C8_with_the_folder_unreachable_the_architecture_is_received_from_its_fire_alarm_ifc(p, monkeypatch, tmp_path):
    _fa_in_force(monkeypatch, _fa_dxf(tmp_path / "fa.dxf"))
    _scan(p)
    shutil.move(str(p.root), str(p.root) + "-offline")
    view = _view(p)
    arch = _badge(view, "ARCH")
    assert arch["badge"] == "received_fa_only" and arch["status"] == "available" and arch["read"] is True
    assert _badge(view, "SM")["badge"] == "unreachable"


def test_C8_schedule_workbooks_are_hashed_into_the_published_identity(complete):
    p = complete
    book = p.ifc / "Mechanical" / "MECHANICAL SCHEDULE" / "Schedule of Fans.xlsx"
    book.parent.mkdir()
    wb = openpyxl.Workbook()
    wb.active.append(["TAG", "LOCATION"])
    wb.active.append(["SPF-01", "ROOF"])
    wb.save(book)
    _scan(p)
    entry = _entry(p, "Schedule of Fans.xlsx")
    assert entry["status"] == "read" and len(entry["sha256"]) == 64
    digest = _row(p).published["sources_digest"]
    wb.active.append(["SPF-02", "ROOF"])
    wb.save(book)
    _scan(p)
    assert _row(p).published["sources_digest"] != digest


# --- T-07..T-11 / C3: revisions, PDFs, confirmed removals -----------------------------------------------


def test_T07_R1_replaced_by_R2_is_retired_and_counted_once(complete):
    p = complete
    ff = p.ifc / "Mechanical/FF"
    shutil.copy(ff / "FF LAYOUT R1.dxf", ff / "FF LAYOUT R2.dxf")
    _rewrite(ff / "FF LAYOUT R2.dxf", "R2")
    (ff / "FF LAYOUT R1.dxf").unlink()
    out = _scan(p)
    assert _entry(p, "FF LAYOUT R2.dxf")["status"] == "read"
    r1 = _entry(p, "FF LAYOUT R1.dxf")
    assert r1["status"] == "removed" and r1["retired_by"].endswith("FF LAYOUT R2.dxf")
    assert out["published"] is True
    view = _view(p)
    assert view["view_state"] == "current"
    assert {r["source"] for r in view["rows"] if r["key"] == "zone_control_valve"} == {"FF LAYOUT R2.dxf"}


def test_T08_T09_an_older_revision_is_superseded_and_read_again_when_the_newer_goes(complete):
    p = complete
    ff = p.ifc / "Mechanical/FF"
    shutil.copy(ff / "FF LAYOUT R1.dxf", ff / "FF LAYOUT R2.dxf")
    _rewrite(ff / "FF LAYOUT R2.dxf", "R2")
    _scan(p)
    r1 = _entry(p, "FF LAYOUT R1.dxf")
    assert r1["status"] == "superseded" and r1["retired_by"].endswith("FF LAYOUT R2.dxf")
    (ff / "FF LAYOUT R2.dxf").unlink()
    _scan(p)
    assert _entry(p, "FF LAYOUT R1.dxf")["status"] == "read"
    r2 = _entry(p, "FF LAYOUT R2.dxf")
    assert (r2["status"], r2["stale_reason"]) == ("stale", "missing")      # an older one is no successor (C1)


def test_C3_a_newer_revision_not_read_does_not_retire_the_older(complete, monkeypatch):
    p = complete
    settings.fa_read_cloud_only_files = False
    try:
        ff = p.ifc / "Mechanical/FF"
        shutil.copy(ff / "FF LAYOUT R1.dxf", ff / "FF LAYOUT R2.dxf")
        _rewrite(ff / "FF LAYOUT R2.dxf", "R2")
        monkeypatch.setattr(evidence, "file_attributes",
                            lambda path, st: evidence.ATTR_RECALL_ON_DATA_ACCESS if "R2" in path else 0)
        _scan(p)
    finally:
        settings.fa_read_cloud_only_files = True
    r1, r2 = _entry(p, "FF LAYOUT R1.dxf"), _entry(p, "FF LAYOUT R2.dxf")
    assert r1["status"] == "superseded" and r1["retired_by"] is None
    assert r2["status"] == "unread" and r2["stale_reason"] == "not_synced"
    view = _view(p)
    assert view["primary"] == "published"                                     # R1 was published; nothing current replaces it


def test_T10_T11_a_drawing_replaced_by_a_pdf_needs_the_engineers_confirmation(complete):
    p = complete
    hvac = p.ifc / "Mechanical/HVAC"
    (hvac / "GROUND FLOOR VENTILATION LAYOUT.dxf").unlink()
    (hvac / "GROUND FLOOR VENTILATION LAYOUT.pdf").write_bytes(b"%PDF-1.4")
    _scan(p)
    assert _entry(p, "VENTILATION LAYOUT.dxf")["stale_reason"] == "missing"
    assert _entry(p, "VENTILATION LAYOUT.pdf")["status"] == "unsupported"
    view = _view(p)
    assert _badge(view, "HVAC")["badge"] == "received_unreadable" and view["primary"] == "published"
    gone = _entry(p, "VENTILATION LAYOUT.dxf")["relative_path"]
    after = p.client.post(f"/projects/{p.pid}/fa-interfaces/sources/confirm-removed", json={"relative_paths": [gone]})
    assert after.status_code == 200 and _entry(p, "VENTILATION LAYOUT.dxf")["status"] == "removed"
    out = _scan(p)                                                            # the next complete read publishes
    assert out["published"] is True and _view(p)["view_state"] == "current"
    assert not any(r["key"] == "ahu" for r in _view(p)["rows"])


def test_C5_sync_states_cannot_be_confirmed_away(complete):
    p = complete
    shutil.move(str(p.ifc), str(p.ifc) + "-cloud")
    _scan(p)
    path = _entry(p, "SM LAYOUT.dxf")["relative_path"]
    refused = p.client.post(f"/projects/{p.pid}/fa-interfaces/sources/confirm-removed", json={"relative_paths": [path]})
    assert refused.status_code == 422 and _entry(p, "SM LAYOUT.dxf")["status"] == "stale"


# --- T-12 / C5: the engineer's way out ------------------------------------------------------------------


def test_T12_publish_current_with_a_reason_the_digest_shown_and_the_compare_and_set(complete, monkeypatch):
    p = complete
    _rewrite(p.ifc / "Mechanical/FF" / "FF LAYOUT R1.dxf")
    _failing_read(monkeypatch, "FF LAYOUT")
    _scan(p)
    view = _view(p)
    assert view["primary"] == "published"
    url = f"/projects/{p.pid}/fa-interfaces/publish-current"
    assert p.client.post(url, json={"reason": "FF drawing is corrupt; contractor reissuing",
                                    "expected_sources_digest": "0" * 64}).status_code == 409
    generation = _row(p).generation
    after = p.client.post(url, json={"reason": "FF drawing is corrupt; contractor reissuing",
                                     "expected_sources_digest": view["current_sources_digest"]})
    assert after.status_code == 200, after.text
    row = _row(p)
    assert (row.published_basis, row.published_reason, row.generation) == (
        "engineer_accepted", "FF drawing is corrupt; contractor reissuing", generation + 1)
    assert row.published_by_id is not None
    assert not any(s["relative_path"].endswith("FF LAYOUT R1.dxf") for s in row.published["sources"])
    shown = after.json()
    assert shown["primary"] == "current" and shown["view_state"] == "provisional"   # FF still not current: said


def test_C5_publish_refuses_an_empty_set_unless_on_purpose(complete):
    p = complete
    for sub in ("SM", "FF", "HVAC"):
        shutil.rmtree(p.ifc / "Mechanical" / sub)
    _scan(p)
    view = _view(p)
    body = {"reason": "the trades withdrew their drawings", "expected_sources_digest": view["current_sources_digest"]}
    assert p.client.post(f"/projects/{p.pid}/fa-interfaces/publish-current", json=body).status_code == 422
    assert p.client.post(f"/projects/{p.pid}/fa-interfaces/publish-current",
                         json={**body, "override": True}).status_code == 200


def test_C5_a_confirmation_during_a_read_makes_the_read_fail_and_stands(complete):
    p = complete
    hvac = p.ifc / "Mechanical/HVAC" / "GROUND FLOOR VENTILATION LAYOUT.dxf"
    hvac.unlink()
    (hvac.parent / "notes.txt").write_text("x")
    _scan(p)
    path = _entry(p, "VENTILATION LAYOUT.dxf")["relative_path"]
    _rewrite(p.ifc / "Mechanical/FF" / "FF LAYOUT R1.dxf")
    fired = []

    def progress(done, total, message, file):
        if not fired:
            fired.append(1)
            from app.database import SessionLocal

            other = SessionLocal()
            try:
                service.confirm_removed(other, other.get(Project, p.pid), 1, [path])
            finally:
                other.close()

    with pytest.raises(service.SourcesChanged):
        _scan(p, progress=progress)
    assert _entry(p, "VENTILATION LAYOUT.dxf")["status"] == "removed"        # the confirmation stands


# --- T-13: live decisions on the published view -------------------------------------------------------


def test_T13_a_decision_made_while_the_published_schedule_is_shown_applies_to_it(complete):
    p = complete
    shutil.move(str(p.ifc / "Mechanical/SM"), str(p.ifc / "Mechanical/SM-moved"))
    _scan(p)
    view = _view(p)
    assert view["primary"] == "published"
    item = next(g for g in view["verification"] if g["key"] == "motorized_smoke_fire_damper")
    after = p.client.post(f"/projects/{p.pid}/fa-interfaces/decisions",
                          json={"id": item["id"], "action": "resolve", "floor_keys": ["L3"], "qty": 1}).json()
    assert after["primary"] == "published"
    assert [r["floor_key"] for r in after["rows"] if r["key"] == "motorized_smoke_fire_damper"] == ["L3"]


# --- T-14: the seeding migration ---------------------------------------------------------------------


def test_T14_the_migration_seeds_once_from_saved_readings_without_a_generation(tmp_path):
    from alembic import command
    from alembic.config import Config
    from sqlalchemy import create_engine, text

    from app.migrations import ALEMBIC_INI

    engine = create_engine(f"sqlite:///{tmp_path / 'm.db'}")
    config = Config(str(ALEMBIC_INI))
    with engine.begin() as conn:
        config.attributes["connection"] = conn
        command.upgrade(config, "c5e7a9b1d3f5")
        def insert(table: str, **values):
            # every NOT NULL column without a default gets a plain value
            for _cid, name, ctype, notnull, default, pk in conn.execute(text(f"PRAGMA table_info({table})")):
                if notnull and default is None and not pk and name not in values:
                    values[name] = 0 if "INT" in (ctype or "").upper() or "BOOL" in (ctype or "").upper() else (
                        "2026-01-01" if "DATE" in (ctype or "").upper() else "x")
            cols = ", ".join(values)
            conn.execute(text(f"INSERT INTO {table} ({cols}) VALUES ({', '.join(':' + c for c in values)})"), values)

        insert("users", id=1, email="m@x", role="admin")
        insert("projects", id=1, ep_number="1", created_by_id=1)
        insert("projects", id=2, ep_number="2", created_by_id=1)
        read = [{"kind": "folder", "discipline": "SM", "relative_path": "x.dxf", "sha256": "a" * 64, "status": "read",
                 "result": {"scan_version": "3", "items": []}},
                {"kind": "folder", "discipline": "FF", "relative_path": "y.dxf", "status": "failed"}]
        conn.execute(text("INSERT INTO project_fa_interfaces (project_id, sources, decisions, manual, updated_at, scanned_at) "
                          "VALUES (1, :s, '{}', '[]', '2026-01-01', '2026-02-02'), (2, :f, '{}', '[]', '2026-01-01', NULL)"),
                     {"s": json.dumps(read), "f": json.dumps([read[1]])})
        command.upgrade(config, "head")
        rows = {r[0]: r for r in conn.execute(text(
            "SELECT project_id, published, published_basis, generation FROM project_fa_interfaces"))}
    seeded = json.loads(rows[1][1])
    assert rows[1][2] == "seeded" and rows[1][3] == 0
    assert [s["relative_path"] for s in seeded["sources"]] == ["x.dxf"]
    assert seeded["sources_digest"] == evidence.digest(seeded["sources"])    # the frozen copy matches the code
    assert rows[2][1] is None and rows[2][3] == 0


# --- T-15 / T-16: concurrent writers, cancel ----------------------------------------------------------


def test_T15_two_reads_racing_the_second_fails_and_generation_moves_once(complete):
    p = complete
    _rewrite(p.ifc / "Mechanical/FF" / "FF LAYOUT R1.dxf")
    raced = []

    def progress(done, total, message, file):
        if not raced:
            raced.append(1)
            from app.database import SessionLocal

            other = SessionLocal()
            try:
                service.scan_project(other, other.get(Project, p.pid))
            finally:
                other.close()

    generation = _row(p).generation
    with pytest.raises(service.SourcesChanged):
        _scan(p, progress=progress)
    assert _row(p).generation == generation + 1


def test_T16_a_cancelled_read_leaves_everything_as_it_was(complete):
    p = complete
    row = _row(p)
    state = (json.dumps(row.sources, sort_keys=True), row.generation, json.dumps(row.published, sort_keys=True))
    _rewrite(p.ifc / "Mechanical/FF" / "FF LAYOUT R1.dxf")
    calls = []

    def check():
        calls.append(1)
        if len(calls) > 1:
            raise jobs.Cancelled()

    with pytest.raises(jobs.Cancelled):
        _scan(p, check=check)
    row = _row(p)
    assert (json.dumps(row.sources, sort_keys=True), row.generation, json.dumps(row.published, sort_keys=True)) == state


# --- T-17 / T-18 / C6: OneDrive cloud-only files -----------------------------------------------------


def _cloud(monkeypatch, fragment: str):
    monkeypatch.setattr(evidence, "file_attributes",
                        lambda path, st: evidence.ATTR_RECALL_ON_DATA_ACCESS if fragment in path else 0)


def test_T17_a_cloud_only_file_that_cannot_come_down_is_not_synced_then_read(complete, monkeypatch):
    p = complete
    _rewrite(p.ifc / "Mechanical/FF" / "FF LAYOUT R1.dxf")
    _cloud(monkeypatch, "FF LAYOUT")

    def sha(path):
        if "FF LAYOUT" in str(path):
            raise OSError("The cloud file provider is not running")
        return REAL_SHA(path)
    monkeypatch.setattr(service, "_sha256", sha)
    _scan(p)
    ff = _entry(p, "FF LAYOUT R1.dxf")
    assert (ff["status"], ff["stale_reason"]) == ("stale", "not_synced") and ff["last_known"]
    assert _view(p)["primary"] == "published"
    monkeypatch.setattr(service, "_sha256", REAL_SHA)
    _scan(p)
    assert _entry(p, "FF LAYOUT R1.dxf")["status"] == "read"


def test_T18_cloud_only_files_left_alone_until_download_and_read(complete, monkeypatch):
    p = complete
    _rewrite(p.ifc / "Mechanical/FF" / "FF LAYOUT R1.dxf")
    _cloud(monkeypatch, "FF LAYOUT")
    opened = []
    monkeypatch.setattr(service, "_sha256", lambda path: opened.append(str(path)) or REAL_SHA(path))
    settings.fa_read_cloud_only_files = False
    try:
        _scan(p)
        assert not any("FF LAYOUT" in o for o in opened)                     # not opened
        assert _entry(p, "FF LAYOUT R1.dxf")["stale_reason"] == "not_synced"
        _scan(p, hydrate=True)                                              # "Download and read"
        assert _entry(p, "FF LAYOUT R1.dxf")["status"] == "read"
    finally:
        settings.fa_read_cloud_only_files = True


def test_C6_a_dehydrated_unchanged_project_rescans_to_current_without_opening_a_file(complete, monkeypatch):
    p = complete
    _cloud(monkeypatch, "")                                                   # every file cloud-only
    opened = []
    monkeypatch.setattr(service, "_sha256", lambda path: opened.append(path) or REAL_SHA(path))
    out = _scan(p)
    assert opened == [] and out["read"] == 3
    view = _view(p)
    assert view["view_state"] == "current" and view["evidence"]["cloud_only"] == 3
    assert {f["status"] for c in view["coverage"] for f in c["files"]} == {"read"}


# --- T-20 / T-22 / T-23 / T-24 / T-27 ----------------------------------------------------------------


def test_T20_the_same_bytes_in_two_packages_are_both_read_and_counted_as_today(p):
    shutil.copy(p.ifc / "Mechanical/HVAC" / "GROUND FLOOR VENTILATION LAYOUT.dxf",
                p.ifc / "Mechanical/SM" / "GROUND FLOOR VENTILATION LAYOUT.dxf")
    _scan(p)
    copies = [e for e in _row(p).sources if e["filename"] == "GROUND FLOOR VENTILATION LAYOUT.dxf"]
    assert [e["status"] for e in copies] == ["read", "read"]
    assert sorted(bool(e.get("duplicate_of")) for e in copies) == [False, True]
    rows = [r for r in _view(p)["rows"] if r["key"] == "ahu"]
    assert len(rows) == 1                                                    # one AHU, as today


def test_T22_a_project_never_read(p):
    view = _view(p)
    assert (view["view_state"], view["primary"], view["rows"], view["totals_known"]) == ("not_read", "none", [], False)


def test_C2_nothing_read_and_no_published_schedule_shows_no_totals_and_publishes_nothing(p, monkeypatch):
    _cloud(monkeypatch, "")
    settings.fa_read_cloud_only_files = False
    try:
        out = _scan(p)
    finally:
        settings.fa_read_cloud_only_files = True
    assert out["read"] == 0 and out["published"] is False and _row(p).published is None
    view = _view(p)
    assert (view["primary"], view["rows"], view["totals_known"]) == ("none", [], False)


def test_T23_received_counts_only_files_there_now(complete):
    p = complete
    shutil.move(str(p.ifc / "Mechanical/SM"), str(p.ifc / "Mechanical/SM-moved"))
    _scan(p)
    view = _view(p)
    available = {c["discipline"] for c in view["coverage"] if c["status"] == "available" and c["folder"]}
    assert available == {c["discipline"] for c in view["coverage"] if c["folder"] and c["received"]}
    assert "SM" not in available and {"FF", "HVAC"} <= available


def test_T24_a_file_changed_after_the_read_is_not_current_on_the_page(complete, monkeypatch):
    p = complete
    path = p.ifc / "Mechanical/HVAC" / "GROUND FLOOR VENTILATION LAYOUT.dxf"
    os.utime(path, (path.stat().st_atime, path.stat().st_mtime + 100))
    view = _view(p)
    assert view["primary"] == "published" and view["evidence"]["changed_since_read"] == 1
    assert any(k["reason"] == "changed_since_read" for k in view["last_known"])
    # an unchanged file that went cloud-only stays current
    os.utime(path, (path.stat().st_atime, path.stat().st_mtime - 100))
    _cloud(monkeypatch, "VENTILATION")
    view = _view(p)
    assert view["view_state"] == "current"
    hvac = next(f for f in _badge(view, "HVAC")["files"] if f["filename"].startswith("GROUND"))
    assert hvac["status"] == "read" and hvac["cloud_only"] is True


def test_T27_the_ignore_list(p):
    sm = p.ifc / "Mechanical/SM"
    for name in ("desktop.ini", "Thumbs.db", "~$SM LAYOUT.dwg", "SM LAYOUT.dwl", "SM LAYOUT.dwl2"):
        (sm / name).write_bytes(b"x")
    _scan(p)
    names = {e["filename"] for e in _row(p).sources}
    assert not names & {"desktop.ini", "Thumbs.db", "~$SM LAYOUT.dwg", "SM LAYOUT.dwl", "SM LAYOUT.dwl2"}


# --- T-25 / C4 / C7: the redesign and the draftsman check --------------------------------------------


def test_C4_kept_interface_changes_are_restored_byte_for_byte():
    from app.redesign import service as R

    approved = {"id": "if-1", "source": R.INTERFACE, "status": "approved", "moved": True, "insert": [1.0, 2.0],
                "interface": {"row": "SM|x|...", "code": "CR"}, "verified": True}
    review = {"id": "rv-1", "source": "review", "status": "pending", "insert": [5.0, 5.0]}
    kept = R.keep_interfaces([approved, review])
    assert list(kept) == ["if-1"]
    changes = [dict(review), json.loads(json.dumps(approved))]
    changes[1]["insert"] = [9.9, 9.9]                                       # what coordinating could do
    changes[1]["status"] = "proposed"
    out = R.restore_kept(changes, kept)
    assert json.dumps(out[1], sort_keys=True) == json.dumps(approved, sort_keys=True)
    assert out[0] == review


def test_T25_add_interfaces_refuses_an_unverified_schedule_and_changes_nothing(complete, monkeypatch):
    from app.models import ProjectIfcDrawing
    from app.redesign import service as R

    p = complete
    drawing = ProjectIfcDrawing(project_id=p.pid, filename="FA.dwg", stored_path="x/FA.dwg")
    p.db.add(drawing)
    p.db.commit()
    row = R.state(p.db, _project(p), drawing.id)
    row.changes = [{"id": "if-1", "source": R.INTERFACE, "status": "approved", "interface": {"row": "a"}}]
    p.db.commit()
    shutil.move(str(p.ifc / "Mechanical/SM"), str(p.ifc / "Mechanical/SM-moved"))
    _scan(p)
    with pytest.raises(R.RedesignError, match="not verified now"):
        R.add_interfaces(p.db, _project(p), drawing.id)
    p.db.expire_all()
    assert R.state(p.db, _project(p), drawing.id).changes == [
        {"id": "if-1", "source": R.INTERFACE, "status": "approved", "interface": {"row": "a"}}]


def test_C7_the_draftsman_check_passes_only_on_a_current_primary_schedule(complete):
    from app.services import draftsman_assignment as D

    p = complete
    ok, note = D._interfaces(p.db, _project(p))
    assert ok and "interfaces" in note
    shutil.move(str(p.ifc / "Mechanical/SM"), str(p.ifc / "Mechanical/SM-moved"))
    _scan(p)
    ok, note = D._interfaces(p.db, _project(p))
    assert not ok and "last published, not verified now" in note


# --- T-26: exports ---------------------------------------------------------------------------------------


def test_T26_exports_say_what_the_schedule_is_and_keep_stale_apart(complete):
    p = complete
    wb = openpyxl.load_workbook(io.BytesIO(p.client.get(f"/projects/{p.pid}/fa-interfaces/export.xlsx").content))
    assert "Last known, not current" not in wb.sheetnames
    assert wb["F. Totals"]["A2"].value.startswith("Current: every drawing verified now")
    shutil.move(str(p.ifc / "Mechanical/SM"), str(p.ifc / "Mechanical/SM-moved"))
    _scan(p)
    wb = openpyxl.load_workbook(io.BytesIO(p.client.get(f"/projects/{p.pid}/fa-interfaces/export.xlsx").content))
    assert wb["C. Interface Schedule"]["A2"].value.startswith("LAST PUBLISHED")
    stale = wb["Last known, not current"]
    assert any("SM LAYOUT.dxf" in str(c.value) for row in stale.iter_rows() for c in row)
    pdf = p.client.get(f"/projects/{p.pid}/fa-interfaces/export.pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")


def test_T26_with_no_schedule_the_exports_print_no_numbers(p, monkeypatch):
    _cloud(monkeypatch, "")
    settings.fa_read_cloud_only_files = False
    try:
        _scan(p)
    finally:
        settings.fa_read_cloud_only_files = True
    wb = openpyxl.load_workbook(io.BytesIO(p.client.get(f"/projects/{p.pid}/fa-interfaces/export.xlsx").content))
    totals = {r[0]: r[1] for r in wb["F. Totals"].iter_rows(min_row=5, values_only=True) if r[0]}
    assert totals["Interface lines"] == "—" and wb["F. Totals"]["A2"].value.startswith("NOT READ COMPLETELY")
